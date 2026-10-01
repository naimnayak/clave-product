"""Word (.docx) export of a saved resume: one clean, single-column document that ATS parsers and Word both read well.

The web app resolves the template-specific bits (headings, job role, accent, font) and sends them along,
so the document uses exactly the labels shown in the editor preview.
"""

import io
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

PAGE_WIDTH_MM = 210
MARGIN_MM = 16
TEXT_WIDTH_MM = PAGE_WIDTH_MM - 2 * MARGIN_MM

DEFAULT_HEADINGS = {
    "summary": "Summary",
    "experience": "Experience",
    "education": "Education",
    "projects": "Projects",
    "skills": "Skills",
    "certifications": "Certifications",
}
DEFAULT_SKILL_LABELS = {"technical": "Technical Skills", "tools": "Tools", "other": "Other Relevant Skills"}
DEFAULT_ORDER = ["experience", "projects", "education", "skills", "certifications"]


def _rgb(hex_color: str) -> RGBColor:
    value = hex_color.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _clean(value: Any) -> str:
    return " ".join(str(value or "").split())


def _bottom_border(paragraph, hex_color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), hex_color.lstrip("#"))
    borders.append(bottom)
    p_pr.append(borders)


class _Writer:
    def __init__(self, *, accent: str, font: str, align_center: bool, uppercase: bool):
        self.accent = accent
        self.uppercase = uppercase
        self.align = WD_ALIGN_PARAGRAPH.CENTER if align_center else WD_ALIGN_PARAGRAPH.LEFT
        self.doc = Document()
        section = self.doc.sections[0]
        section.page_width, section.page_height = Mm(PAGE_WIDTH_MM), Mm(297)
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Mm(MARGIN_MM))
        section.top_margin = section.bottom_margin = Mm(14)
        normal = self.doc.styles["Normal"]
        normal.font.name = font
        normal.font.size = Pt(10.5)
        # East Asian font slot too, or Word substitutes its own default for some text.
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.line_spacing = 1.1

    def para(self, text: str = "", *, size: float | None = None, bold=False, italic=False, color: str | None = None, align=None, space_before=0.0, space_after=0.0):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(space_before)
        paragraph.paragraph_format.space_after = Pt(space_after)
        if align is not None:
            paragraph.alignment = align
        if text:
            self.run(paragraph, text, size=size, bold=bold, italic=italic, color=color)
        return paragraph

    def run(self, paragraph, text: str, *, size: float | None = None, bold=False, italic=False, color: str | None = None):
        run = paragraph.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = _rgb(color)
        return run

    def split_line(self, left: str, right: str, *, bold_left=True, space_before=0.0):
        """Left text with right-aligned text on the same line (title + dates)."""
        paragraph = self.para(space_before=space_before)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT)
        if left:
            self.run(paragraph, left, bold=bold_left)
        if right:
            self.run(paragraph, f"\t{right}")
        return paragraph

    def heading(self, text: str) -> None:
        paragraph = self.para(text.upper() if self.uppercase else text, size=11, bold=True, color=self.accent, space_before=10, space_after=4)
        paragraph.paragraph_format.keep_with_next = True
        _bottom_border(paragraph, self.accent)

    def bullets(self, items: list[str]) -> None:
        for item in items:
            text = _clean(item)
            if text:
                paragraph = self.doc.add_paragraph(text, style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(1)

    def save(self) -> bytes:
        buffer = io.BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()


def build_resume_docx(resume: dict[str, Any], options: dict[str, Any]) -> bytes:
    content = resume.get("content") or {}
    contact = content.get("contact") or {}
    headings = {**DEFAULT_HEADINGS, **{k: v for k, v in (content.get("headings") or {}).items() if _clean(v)}, **{k: v for k, v in (options.get("headings") or {}).items() if _clean(v)}}
    skill_labels = {**DEFAULT_SKILL_LABELS, **{k: v for k, v in (content.get("skillLabels") or {}).items() if _clean(v)}, **{k: v for k, v in (options.get("skillLabels") or {}).items() if _clean(v)}}

    w = _Writer(
        accent=options.get("accent") or "#056B4D",
        font=options.get("font") or "Calibri",
        align_center=options.get("align") == "center",
        uppercase=bool(options.get("uppercase", True)),
    )

    # ── Header: name, job role, tagline, contact line ──
    w.para(_clean(contact.get("name")) or "Your Name", size=22, bold=True, align=w.align)
    role = _clean(options.get("role"))
    if role:
        w.para(role, size=12, color=w.accent, align=w.align, space_before=2)
    tagline = _clean(contact.get("tagline"))
    if tagline:
        w.para(tagline, size=10.5, color="#555555", align=w.align, space_before=2)
    contact_line = " | ".join(_clean(contact.get(key)) for key in ("email", "phone", "location", "linkedin", "github", "portfolio") if _clean(contact.get(key)))
    if contact_line:
        w.para(contact_line, size=9.5, color="#444444", align=w.align, space_before=4)

    summary = _clean(content.get("summary"))
    if summary:
        w.heading(headings["summary"])
        w.para(summary)

    def experience() -> None:
        items = content.get("experience") or []
        if not items:
            return
        w.heading(headings["experience"])
        for index, item in enumerate(items):
            dates = " – ".join(part for part in (_clean(item.get("start")), _clean(item.get("end"))) if part)
            w.split_line(_clean(item.get("title")), dates, space_before=0 if index == 0 else 6)
            where = ", ".join(part for part in (_clean(item.get("company")), _clean(item.get("location"))) if part)
            if where:
                w.para(where, italic=True)
            if _clean(item.get("description")):
                w.para(_clean(item.get("description")), space_before=1)
            w.bullets(item.get("bullets") or [])

    def education() -> None:
        items = content.get("education") or []
        if not items:
            return
        w.heading(headings["education"])
        for index, item in enumerate(items):
            degree, institution = _clean(item.get("degree")), _clean(item.get("institution"))
            w.split_line(degree or institution, _clean(item.get("dates")), space_before=0 if index == 0 else 6)
            where = ", ".join(part for part in (institution, _clean(item.get("location"))) if part)
            if degree and where:
                w.para(where)
            if _clean(item.get("details")):
                w.para(_clean(item.get("details")))

    def projects() -> None:
        items = content.get("projects") or []
        if not items:
            return
        w.heading(headings["projects"])
        for index, item in enumerate(items):
            w.split_line(_clean(item.get("name")), _clean(item.get("link")), space_before=0 if index == 0 else 6)
            tech = ", ".join(_clean(t) for t in item.get("tech") or [] if _clean(t))
            if tech:
                w.para(tech, italic=True)
            if _clean(item.get("description")):
                w.para(_clean(item.get("description")), space_before=1)
            w.bullets(item.get("bullets") or [])

    def skills() -> None:
        groups = content.get("skills") or {}
        filled = [(key, [_clean(s) for s in groups.get(key) or [] if _clean(s)]) for key in ("technical", "tools", "other")]
        filled = [(key, values) for key, values in filled if values]
        if not filled:
            return
        w.heading(headings["skills"])
        if options.get("skillsMode") == "chips":
            w.para(" • ".join(value for _, values in filled for value in values))
            return
        for key, values in filled:
            paragraph = w.para()
            w.run(paragraph, f"{_clean(skill_labels[key])}: ", bold=True)
            w.run(paragraph, ", ".join(values))

    def certifications() -> None:
        items = content.get("certifications") or []
        if not items:
            return
        w.heading(headings["certifications"])
        for item in items:
            paragraph = w.split_line(_clean(item.get("name")), _clean(item.get("date")))
            # Issuer and link sit after the bold name, before the right-aligned date.
            extra = (f", {_clean(item.get('issuer'))}" if _clean(item.get("issuer")) else "") + (f" ({_clean(item.get('link'))})" if _clean(item.get("link")) else "")
            if extra:
                tab_run = paragraph.runs[-1] if _clean(item.get("date")) else None
                new_run = paragraph.add_run(extra)
                if tab_run is not None:
                    tab_run._r.addprevious(new_run._r)

    writers = {"experience": experience, "education": education, "projects": projects, "skills": skills, "certifications": certifications}
    # Same order as the preview, which only renders the sections listed in sectionOrder.
    order = list(dict.fromkeys(resume.get("sectionOrder") or DEFAULT_ORDER))
    for key in order:
        if key in writers:
            writers[key]()
    return w.save()
